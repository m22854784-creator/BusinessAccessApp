"""Report builders plus PDF (reportlab) and Excel (openpyxl) renderers."""
from datetime import timedelta, timezone
from io import BytesIO
from xml.sax.saxutils import escape

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from db import parse_iso, utcnow

BRAND = "0A6CB5"
MAX_ROWS = 10000
LOGIN_ACTIONS = ("LOGIN", "LOGOUT", "LOGIN_FAILED", "LOGIN_BLOCKED", "ACCOUNT_LOCKED", "SESSION_TIMEOUT")


def local_tz():
    return timezone(timedelta(minutes=current_app.config["TZ_OFFSET_MIN"]))


def tz_label():
    m = current_app.config["TZ_OFFSET_MIN"]
    sign = "+" if m >= 0 else "-"
    return f"UTC{sign}{abs(m) // 60:02d}:{abs(m) % 60:02d}"


def fmt_ts(value):
    if not value:
        return ""
    return parse_iso(value).astimezone(local_tz()).strftime("%d-%b-%Y %H:%M:%S")


def _module_names():
    return {m["key"]: m["name"] for m in current_app.config["MODULES"]}


def build_report(db, name, module=None):
    """Return {title, headers, rows, weights} for a named report."""
    modules = current_app.config["MODULES"]
    names = _module_names()

    if name == "audit":
        rows = db.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT ?", (MAX_ROWS,)).fetchall()
        return {
            "title": "Audit Trail",
            "headers": ["Time", "User", "Category", "Action", "Target", "Details", "IP", "Result"],
            "weights": [13, 11, 8, 13, 16, 27, 9, 7],
            "rows": [[fmt_ts(r["ts"]), r["username"] or "-", r["category"], r["action"],
                      r["target"] or "", r["details"] or "", r["ip"] or "",
                      "Success" if r["success"] else "Failed"] for r in rows],
        }

    if name == "logins":
        marks = ",".join("?" * len(LOGIN_ACTIONS))
        rows = db.execute(f"SELECT * FROM audit_log WHERE action IN ({marks}) ORDER BY id DESC LIMIT ?",
                          (*LOGIN_ACTIONS, MAX_ROWS)).fetchall()
        return {
            "title": "Login Activity",
            "headers": ["Time", "User", "Event", "Result", "IP address", "Details"],
            "weights": [15, 13, 15, 8, 12, 37],
            "rows": [[fmt_ts(r["ts"]), r["username"] or "-", r["action"],
                      "Success" if r["success"] else "Failed", r["ip"] or "", r["details"] or ""]
                     for r in rows],
        }

    if name == "users":
        perms = {}
        for p in db.execute("SELECT user_id, module, level FROM permissions"):
            perms.setdefault(p["user_id"], {})[p["module"]] = p["level"]
        rows = []
        for u in db.execute("SELECT * FROM users ORDER BY role, username"):
            def lvl(key, u=u):
                if u["role"] == "admin":
                    return "Edit (admin)"
                return {"none": "No Access", "view": "View", "edit": "Edit"}[perms.get(u["id"], {}).get(key, "none")]
            rows.append([u["username"], u["full_name"], u["role"].title(), u["department"] or "",
                         "Active" if u["is_active"] else "Disabled", fmt_ts(u["last_login"]) or "Never",
                         u["login_count"], *[lvl(m["key"]) for m in modules]])
        return {
            "title": "User & Permission Register",
            "headers": ["Username", "Name", "Role", "Department", "Status", "Last login", "Logins",
                        *[m["name"] for m in modules]],
            "weights": [11, 13, 6, 9, 7, 13, 5, 9, 9, 9, 9],
            "rows": rows,
        }

    if name == "records":
        sql, args = "SELECT * FROM records", []
        if module:
            sql, args = sql + " WHERE module=?", [module]
        rows = db.execute(sql + " ORDER BY module, id", args).fetchall()
        title = "Records - " + names[module] if module else "Shop Records - All Modules"
        return {
            "title": title,
            "headers": ["Module", "Part no.", "Description", "Qty", "Status", "Updated by", "Updated at"],
            "weights": [11, 9, 30, 6, 10, 12, 14],
            "rows": [[names.get(r["module"], r["module"]), r["part_no"], r["description"], r["quantity"],
                      r["status"], r["updated_by"] or "", fmt_ts(r["updated_at"])] for r in rows],
        }
    raise KeyError(name)


def _meta(generated_by):
    return f"Generated {utcnow().astimezone(local_tz()).strftime('%d-%b-%Y %H:%M')} ({tz_label()}) by {generated_by}"


def to_xlsx(report, generated_by):
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31]
    ncols = len(report["headers"])
    ws["A1"] = f"Godrej & Boyce - {report['title']}"
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws["A2"] = _meta(generated_by)
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(report["headers"])
    header_row = 4
    for c in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(vertical="center")
    for row in report["rows"]:
        ws.append(row)
    widths = [len(str(h)) for h in report["headers"]]
    for row in report["rows"][:500]:
        for i, v in enumerate(row):
            widths[i] = max(widths[i], min(len(str(v)), 60))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 3
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{max(header_row, ws.max_row)}"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _cell(text, style):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)


def to_pdf(report, generated_by):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=22 * mm, bottomMargin=14 * mm, title=report["title"],
                            author="Godrej & Boyce BAMS")
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=7.5, leading=9.2)
    head = ParagraphStyle("head", parent=body, fontName="Helvetica-Bold", textColor=colors.white)
    title = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#" + BRAND))
    meta = ParagraphStyle("meta", fontName="Helvetica-Oblique", fontSize=8, textColor=colors.HexColor("#666666"))

    data = [[_cell(h, head) for h in report["headers"]]]
    rows = report["rows"] or [["No records to show"] + [""] * (len(report["headers"]) - 1)]
    data += [[_cell(v, body) for v in row] for row in rows]
    total = float(sum(report["weights"]))
    widths = [doc.width * w / total for w in report["weights"]]

    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BRAND)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    def decorate(canvas, d):
        w, h = landscape(A4)
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#" + BRAND))
        canvas.rect(0, h - 12 * mm, w, 12 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawString(12 * mm, h - 8.5 * mm, "godrej")
        canvas.setFont("Helvetica", 9)
        canvas.drawRightString(w - 12 * mm, h - 8 * mm, "Business Access Management System")
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(12 * mm, 8 * mm, "Confidential - for internal use only")
        canvas.drawRightString(w - 12 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build([Paragraph(escape(report["title"]), title), Paragraph(escape(_meta(generated_by)), meta),
               Spacer(1, 4 * mm), table], onFirstPage=decorate, onLaterPages=decorate)
    buf.seek(0)
    return buf
